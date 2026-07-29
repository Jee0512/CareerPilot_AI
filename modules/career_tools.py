"""
career_tools.py
----------------
Supporting career-coach features layered on top of the core analysis:
  - Professional outreach/cover email tailored to the resume + JD
  - A personalized learning roadmap for the candidate's missing skills
  - A plain-text/Markdown report bundling everything for download

Report generation is pure string assembly (no AI call) so "Download Report"
always works instantly, even if the AI service is briefly unavailable.
"""

import io
import re
from datetime import date
from modules.gemini_client import generate_text, generate_json


def generate_outreach_email(resume_text: str, jd_text: str) -> str:
    """A short, professional email a candidate could send to a recruiter/hiring manager."""
    prompt = f"""
Write a short, professional outreach email (not a full cover letter) that a
candidate could send to a recruiter or hiring manager, expressing interest
in the role below and briefly highlighting 2-3 relevant strengths from their
resume. Keep it under 150 words, warm but professional, with a subject line.

RESUME:
\"\"\"{resume_text[:3000]}\"\"\"

JOB DESCRIPTION:
\"\"\"{jd_text[:2000]}\"\"\"

Output plain text: "Subject: ..." on the first line, then a blank line, then the email body.
"""
    return generate_text(prompt, temperature=0.5).strip()


def generate_learning_roadmap(missing_skills: list[str], jd_text: str) -> str:
    """A short, actionable learning plan to close the candidate's skill gaps."""
    if not missing_skills:
        return "No missing skills detected — no roadmap needed. You're already well-aligned with this role!"

    prompt = f"""
A candidate is missing these skills for a target role: {missing_skills}.

JOB DESCRIPTION CONTEXT:
\"\"\"{jd_text[:1500]}\"\"\"

Write a concise, actionable learning roadmap to close these gaps, organized
week by week (2-4 weeks), with 1-2 specific resource types per skill (e.g.
"official docs", "a hands-on project", "a short course") — no need for real
URLs, just resource types. Keep it under 250 words. Use plain text with
"Week 1:", "Week 2:" etc. as section labels.
"""
    return generate_text(prompt, temperature=0.4).strip()


def generate_resume_debate(resume_text: str, jd_text: str) -> dict:
    """
    Two-sided recruiter-perspective debate: reasons the resume might be
    shortlisted vs rejected for this specific role, grounded strictly in
    what's actually in the resume and JD — no invented skills/experience.
    """
    prompt = f"""
You are simulating how a recruiter would debate internally about whether to
shortlist this candidate for this specific role. Base everything ONLY on
what's actually in the resume and job description below — never invent
skills, experience, projects, or certifications.

RESUME:
\"\"\"{resume_text[:4000]}\"\"\"

JOB DESCRIPTION:
\"\"\"{jd_text[:2500]}\"\"\"

Return ONLY valid JSON (no markdown fences, no commentary) with this exact shape:
{{
  "shortlist_reasons": ["specific reason grounded in the resume", "..."],
  "rejection_reasons": ["specific concern grounded in the gap between resume and JD", "..."]
}}
Give 3-5 points per side.
"""
    return generate_json(prompt, temperature=0.4)


def generate_bullet_rewrites(resume_text: str, missing_skills: list[str], jd_text: str) -> list[dict]:
    """
    Suggest rewritten versions of 2-3 existing resume bullet points that
    naturally incorporate the candidate's missing skills, without inventing
    fake experience — only rephrasing/reframing what's already there.
    """
    if not missing_skills:
        return []

    prompt = f"""
A candidate is missing these skills for a target role: {missing_skills}.

Their current resume:
\"\"\"{resume_text[:3000]}\"\"\"

Target job description:
\"\"\"{jd_text[:1500]}\"\"\"

Pick 2-3 EXISTING bullet points from the resume that could be reframed to
naturally highlight adjacent/transferable experience toward the missing
skills, WITHOUT fabricating experience the candidate doesn't have. Only
reframe wording/emphasis — never invent new tools, projects, or claims.

Return ONLY valid JSON (no markdown fences, no commentary) as a list:
[
  {{"original": "the exact original bullet point text", "rewritten": "improved version", "why": "one sentence on what changed and why"}}
]
"""
    from modules.gemini_client import generate_json
    result = generate_json(prompt, temperature=0.4)
    return result if isinstance(result, list) else result.get("suggestions", [])


def generate_tailored_resume(resume_text: str, jd_text: str, matched_skills: list[str], missing_skills: list[str] | None = None) -> str:
    """
    Rewrite the candidate's resume tailored to the target JD WITHOUT changing
    its overall format, layout, or section order — only rewords bullets,
    strengthens action verbs, and reorders skills WITHIN a section for ATS
    visibility. Must not invent companies, titles, degrees, dates, projects,
    certifications, or achievements not in the original.
    """
    prompt = f"""
You are an expert ATS optimization specialist. Rewrite this resume to achieve
the HIGHEST POSSIBLE ATS match score (target 85-95%) for the specific job
description below, while NEVER fabricating experience.

ORIGINAL RESUME:
\"\"\"{resume_text[:4000]}\"\"\"

TARGET JOB DESCRIPTION:
\"\"\"{jd_text[:2500]}\"\"\"

ALREADY-MATCHED SKILLS TO EMPHASIZE: {matched_skills}

JD SKILLS CURRENTLY MISSING FROM THE RESUME — if the candidate's real
experience genuinely supports any of these (e.g. they used a similar tool
or performed a related task), use the EXACT JD terminology to describe it.
Do NOT fabricate skills or experience that has zero basis in the original:
{missing_skills or []}

CRITICAL ATS OPTIMIZATION TECHNIQUES:
1. **Keyword Density**: Use the JD's exact terminology and phrases wherever
   the resume already supports it. ATS systems match exact strings.
2. **Action Verbs**: Replace weak verbs with strong ATS-friendly ones:
   "Led", "Architected", "Optimized", "Engineered", "Implemented",
   "Automated", "Designed", "Delivered", "Spearheaded", "Orchestrated".
3. **Quantify Results**: Where numbers exist, surface them prominently.
4. **Section Relevance**: Reorder bullet points within each section so the
   MOST RELEVANT experience to THIS JD appears FIRST.
5. **Skills Section**: Reorder skills so JD-critical keywords come first.
   Use the JD's exact spelling and formatting.
6. **Contextual Weaving**: Naturally incorporate JD keywords into bullet
   point descriptions where the underlying work genuinely relates.

Rules:
- Do NOT invent companies, job titles, dates, degrees, projects,
  certifications, or achievements not in the original.
- Keep the exact same section order as the original.
- Rewrite bullet points with stronger action verbs and JD terminology.
- Output PLAIN TEXT — section headers in CAPITALS, bullet points with "- ".
- Target: 85-95% ATS keyword match while staying 100% truthful.
"""
    return generate_text(prompt, temperature=0.4).strip()


def build_resume_pdf(resume_text: str) -> bytes:
    """
    Render plain-text resume content into a simple, clean multi-page PDF
    using PyMuPDF (already a project dependency for PDF reading). Paginates
    by rough character count per page rather than precise text-fit
    calculation, which is sufficient for a readable, professional layout.
    """
    import fitz

    doc = fitz.open()
    rect = fitz.Rect(50, 50, 545, 792)
    max_chars_per_page = 3200

    page = doc.new_page()
    buffer = ""
    for line in resume_text.split("\n"):
        candidate = buffer + line + "\n"
        if len(candidate) > max_chars_per_page and buffer:
            page.insert_textbox(rect, buffer, fontsize=10.5, fontname="helv")
            page = doc.new_page()
            buffer = line + "\n"
        else:
            buffer = candidate
    if buffer:
        page.insert_textbox(rect, buffer, fontsize=10.5, fontname="helv")

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def build_resume_docx(resume_text: str) -> bytes:
    """Render plain-text resume content into a simple .docx using python-docx."""
    from docx import Document
    from docx.shared import Pt

    document = Document()
    for line in resume_text.split("\n"):
        stripped = line.strip()
        if not stripped:
            document.add_paragraph("")
            continue
        is_header = stripped.isupper() and len(stripped) > 2
        p = document.add_paragraph()
        run = p.add_run(stripped)
        run.bold = is_header
        run.font.size = Pt(13 if is_header else 11)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def build_section_diffs(lines: list[dict], rewrites: dict) -> list[dict]:
    """
    Group layout lines into sections (by detecting header lines using the
    same heuristic that already excludes headers from rewriting) and build
    a before/after diff per section for the Before-vs-After comparison view.
    Returns [{"name": str, "diff_html": str|None, "original": str, "changed": bool}, ...]
    """
    def _is_header(t: str) -> bool:
        t = t.strip()
        return (t.isupper() and len(t) < 25 and "," not in t) or (t.endswith(":") and len(t) < 25)

    sections = []
    current_name = "Header"
    current_lines = []
    for l in lines:
        if _is_header(l["text"]):
            if current_lines:
                sections.append((current_name, current_lines))
            current_name = l["text"].rstrip(":").title()
            current_lines = []
        else:
            current_lines.append(l)
    if current_lines:
        sections.append((current_name, current_lines))

    result = []
    for name, sec_lines in sections:
        original = "\n".join(l["text"] for l in sec_lines)
        improved = "\n".join(rewrites.get(l["id"], l["text"]) for l in sec_lines)
        changed = any(l["id"] in rewrites for l in sec_lines)
        result.append({
            "name": name,
            "diff_html": build_resume_diff_html(original, improved) if changed else None,
            "original": original,
            "changed": changed,
        })
    return result


def build_resume_diff_html(original: str, improved: str) -> str:
    """
    Word-level diff between original and improved resume text, rendered as
    HTML: additions highlighted green, removals struck through in red.
    Used to satisfy "highlight what changed" without a heavyweight diff lib.
    """
    import difflib
    import html

    orig_words = original.split()
    new_words = improved.split()
    matcher = difflib.SequenceMatcher(None, orig_words, new_words)

    parts = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            parts.append(html.escape(" ".join(new_words[j1:j2])))
        elif tag == "insert":
            added = html.escape(" ".join(new_words[j1:j2]))
            parts.append(f'<span style="background:#EAF7EE;color:#1C8A52;font-weight:600;">{added}</span>')
        elif tag == "delete":
            removed = html.escape(" ".join(orig_words[i1:i2]))
            parts.append(f'<span style="background:#FDECEC;color:#C13333;text-decoration:line-through;">{removed}</span>')
        elif tag == "replace":
            removed = html.escape(" ".join(orig_words[i1:i2]))
            added = html.escape(" ".join(new_words[j1:j2]))
            parts.append(f'<span style="background:#FDECEC;color:#C13333;text-decoration:line-through;">{removed}</span>')
            parts.append(f'<span style="background:#EAF7EE;color:#1C8A52;font-weight:600;">{added}</span>')

    return " ".join(parts)


def _is_rewritable_line(text: str) -> bool:
    """Heuristic filter: only rewrite substantive content lines, never
    headers, contact info, dates, or short labels — keeps names, section
    titles, emails, phone numbers, and URLs untouched. Short skill-list
    lines (e.g. "Python, SQL, AWS") ARE kept — they're high-value targets
    for keyword alignment and were previously wrongly excluded."""
    t = text.strip()
    if len(t) < 3:
        return False
    if re.search(r"\S+@\S+\.\S+", t):
        return False
    if re.match(r"^https?://", t) or "linkedin.com" in t.lower() or "github.com" in t.lower():
        return False
    if re.match(r"^[\d\-\+\(\)\s]{7,}$", t):
        return False
    if t.isupper() and len(t) < 25 and "," not in t:
        return False
    if t.endswith(":") and len(t) < 25:
        return False
    if re.search(r"\(\s*\d{4}\s*[-–]\s*(\d{4}|present)\s*\)", t, re.IGNORECASE):
        return False
    return True


def generate_layout_constrained_rewrites(lines: list[dict], jd_text: str, matched_skills: list[str], missing_skills: list[str] | None = None) -> dict:
    """
    Ask Gemini to rewrite eligible resume lines for ATS/JD fit, each
    constrained to fit back into its ORIGINAL fixed-width space in the PDF.
    Returns {line_id: rewritten_text} — only for lines actually changed.
    """
    candidates = [l for l in lines if l["id"] != 0 and _is_rewritable_line(l["text"])]
    if not candidates:
        return {}

    items = [{"id": l["id"], "text": l["text"], "max_len": len(l["text"]) + max(15, int(len(l["text"]) * 0.35))} for l in candidates]

    prompt = f"""
You are optimizing a resume's wording for this job description, LINE BY
LINE, where each line must fit back into a roughly fixed-width space in the
original PDF layout — you cannot make any line longer than its "max_len".

JOB DESCRIPTION:
\"\"\"{jd_text[:2000]}\"\"\"

SKILLS ALREADY ON THE RESUME TO EMPHASIZE WHERE RELEVANT: {matched_skills}

JD SKILLS CURRENTLY MISSING FROM THE RESUME — if the candidate's real
experience genuinely supports any of these (even if not phrased this way
originally), use the exact JD terminology to surface them truthfully:
{missing_skills or []}

GOAL: Maximize keyword/skill alignment with the job description as
aggressively as possible — this resume will be scored by an ATS on exact
keyword overlap, so use the JD's own terminology wherever the candidate's
real experience genuinely supports it. For SKILLS-list lines
specifically, reorder and rephrase to surface every JD-relevant skill the
candidate already has as prominently as possible.

CRITICAL CONSTRAINT: the rewritten text length in characters MUST NOT
exceed that line's "max_len" — it must physically fit in the same space.
If you cannot meaningfully improve a line within that limit, leave it out
of your response entirely (do not include it).

STRICT TRUTHFULNESS RULE: Never invent skills, experience, projects,
achievements, or certifications not already present or clearly implied by
the line. If a JD requirement has zero support anywhere in the resume, do
not fabricate it — leave that gap as-is rather than lie.

LINES (JSON):
{items}

Return ONLY valid JSON (no markdown fences, no commentary) as an object
mapping id (as a string) to the rewritten text — only include ids you
actually improved:
{{"0": "rewritten line 0", "3": "rewritten line 3"}}
"""
    result = generate_json(prompt, temperature=0.4)
    if not isinstance(result, dict):
        return {}
    return {int(k): v for k, v in result.items() if str(k).isdigit()}


def _map_font(font_name: str, flags: int) -> str:
    """Best-effort match of the original embedded font to PyMuPDF's Base14
    fonts (insert_text can't use arbitrary embedded fonts by name)."""
    name = (font_name or "").lower()
    bold = "bold" in name
    italic = "italic" in name or "oblique" in name

    if "courier" in name or "mono" in name:
        table = {(False, False): "cour", (True, False): "cobo", (False, True): "coit", (True, True): "cobi"}
    elif any(k in name for k in ("times", "serif", "georgia", "garamond", "cambria", "book")):
        table = {(False, False): "tiro", (True, False): "tibo", (False, True): "tiit", (True, True): "tibi"}
    else:
        table = {(False, False): "helv", (True, False): "hebo", (False, True): "heit", (True, True): "hebi"}
    return table[(bold, italic)]


def _color_to_rgb(color_int) -> tuple:
    if not color_int:
        return (0, 0, 0)
    r = (color_int >> 16) & 255
    g = (color_int >> 8) & 255
    b = color_int & 255
    return (r / 255, g / 255, b / 255)


def apply_layout_preserving_edits(pdf_bytes: bytes, lines: list[dict], rewrites: dict) -> bytes:
    """
    Edit the ORIGINAL PDF in place: white-out each rewritten line's exact
    bounding box (redaction) and insert the new text at the same position,
    matching the original font family/size/color as closely as PyMuPDF's
    Base14 fonts allow. Everything else — images, icons, tables, other
    text, layout, spacing, margins — is left completely untouched.
    """
    import fitz

    try:
        doc = fitz.open(stream=io.BytesIO(pdf_bytes), filetype="pdf")
    except Exception as exc:
        raise RuntimeError("Couldn't reopen the original PDF for editing.") from exc

    by_page: dict[int, list[dict]] = {}
    for line in lines:
        if line["id"] in rewrites:
            by_page.setdefault(line["page"], []).append(line)

    for page_num, page_lines in by_page.items():
        page = doc[page_num]
        for line in page_lines:
            page.add_redact_annot(fitz.Rect(line["bbox"]), fill=(1, 1, 1))
        page.apply_redactions()

        for line in page_lines:
            new_text = rewrites[line["id"]]
            fontname = _map_font(line["font"], line["flags"])
            color = _color_to_rgb(line["color"])
            x0, y0, x1, y1 = line["bbox"]
            baseline_y = y1 - (y1 - y0) * 0.2
            try:
                page.insert_text((x0, baseline_y), new_text, fontsize=line["size"], fontname=fontname, color=color)
            except Exception:
                page.insert_text((x0, baseline_y), new_text, fontsize=line["size"], fontname="helv", color=color)

    result = doc.tobytes()
    doc.close()
    return result


def lines_to_text(lines: list[dict], rewrites: dict | None = None) -> str:
    """Reconstruct plain text from layout lines (rewrites applied where present)."""
    rewrites = rewrites or {}
    return "\n".join(rewrites.get(l["id"], l["text"]) for l in lines)


def build_report_markdown(state: dict) -> str:
    """
    Assemble a Markdown report from session state (no AI call). `state` is
    expected to be a dict-like view of st.session_state with the relevant keys.
    """
    lines = [
        "# CareerPilot AI — Career Report",
        f"_Generated {date.today().isoformat()}_",
        "",
    ]

    match = state.get("match_result")
    if match:
        lines += [
            "## Resume Match Score",
            f"**{match['match_score']}%**",
            "",
            f"✅ Matching Skills ({match['matched_pct']}%): "
            + (", ".join(match["matched_skills"]) or "none detected"),
            "",
            f"❌ Missing Skills ({match['missing_pct']}%): "
            + (", ".join(match["missing_skills"]) or "none — great fit!"),
            "",
        ]

    evaluation = state.get("evaluation")
    if evaluation:
        cat = evaluation.get("category_scores", {})
        lines += [
            "## Interview Readiness",
            f"**Overall Score:** {evaluation.get('overall_score', '—')}%",
            f"**Readiness Level:** {evaluation.get('readiness_level', '—')}",
            "",
            "### Category Scores",
        ]
        for k in ("aptitude", "verbal", "reasoning", "technical", "coding"):
            if cat.get(k) is not None:
                lines.append(f"- {k.title()}: {cat[k]}%")
        lines += [
            "",
            "### Strong Areas",
            *[f"- {a}" for a in evaluation.get("strong_areas", [])],
            "",
            "### Weak Areas",
            *[f"- {a}" for a in evaluation.get("weak_areas", [])],
            "",
            "### Recommendations",
            *[f"- {r}" for r in evaluation.get("recommendations", [])],
            "",
        ]

    jobs = state.get("recommended_jobs")
    if jobs:
        lines.append("## Recommended Jobs")
        for job in jobs:
            lines += [
                f"### {job.get('job_title', 'Role')} — {job.get('estimated_match_pct', '—')}% match",
                job.get("reason", ""),
                f"- Skills present: {', '.join(job.get('skills_present', []))}",
                f"- Skills to improve: {', '.join(job.get('skills_missing', []))}",
                "",
            ]

    email = state.get("outreach_email")
    if email:
        lines += ["## Outreach Email Draft", "```", email, "```", ""]

    roadmap = state.get("learning_roadmap")
    if roadmap:
        lines += ["## Learning Roadmap", roadmap, ""]

    resume = state.get("tailored_resume")
    if resume:
        lines += ["## AI-Tailored Resume", "```", resume, "```", ""]

    rewrites = state.get("bullet_rewrites")
    if rewrites:
        lines.append("## Resume Bullet Rewrite Suggestions")
        for r in rewrites:
            lines += [
                f"- **Original:** {r.get('original', '')}",
                f"  **Rewritten:** {r.get('rewritten', '')}",
                f"  _{r.get('why', '')}_",
                "",
            ]

    return "\n".join(lines)

